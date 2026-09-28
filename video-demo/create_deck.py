#!/usr/bin/env python3
"""Create TuntasUMKM PPT with diagrams and visual elements."""
import json
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG = "0a0a0a"
TITLE = "f5f5f5"
SUBTITLE = "a3a3a3"
BODY = "d4d4d4"
BLUE = "3b82f6"
GREEN = "22c55e"
ORANGE = "f97316"
RED = "ef4444"
PURPLE = "8b5cf6"
TEAL = "14b8a6"
AMBER = "f59e0b"
CARD_BG = "141414"
CARD_BORDER = "2a2a2a"

def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)

def add_text(slide, l, t, w, h, text, size=20, color=BODY, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor.from_string(color)
    p.alignment = align
    return txBox

def add_rect(slide, l, t, w, h, fill, text="", tc="f5f5f5", fs=14, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border:
        shape.line.color.rgb = RGBColor.from_string(border)
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(fs)
        p.font.color.rgb = RGBColor.from_string(tc)
        p.alignment = PP_ALIGN.CENTER
    return shape

def add_arrow(slide, l, t, w, h, color="525252"):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape

def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def title_slide(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, spec.get("background", BG))
    add_text(slide, 0.8, 2.0, 11.7, 1.5, spec["title"], 72, TITLE, True, PP_ALIGN.CENTER)
    if spec.get("subtitle"):
        add_text(slide, 0.8, 3.8, 11.7, 2.0, spec["subtitle"], 28, SUBTITLE, False, PP_ALIGN.CENTER)
    if spec.get("notes"):
        add_notes(slide, spec["notes"])
    return slide

def arch_slide(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, spec.get("background", BG))
    add_text(slide, 0.8, 0.3, 11.7, 0.8, spec["title"], 40, TITLE, True)
    if spec.get("subtitle"):
        add_text(slide, 0.8, 1.1, 11.7, 0.5, spec["subtitle"], 20, SUBTITLE)
    
    # Left: vertical flow
    fx, fw, bh, gap = 1.5, 2.0, 0.6, 0.35
    add_rect(slide, fx, 2.0, fw, bh, "1a1a1a", "Customer", "94a3b8", 14, "94a3b8")
    add_arrow(slide, fx+0.85, 2.6, 0.3, gap)
    add_rect(slide, fx, 2.95, fw, bh, "1a1a1a", "Hermes Agent", BLUE, 14, BLUE)
    add_arrow(slide, fx+0.85, 3.55, 0.3, gap)
    add_rect(slide, fx, 3.9, fw, bh, "1a1a1a", "MCP", PURPLE, 14, PURPLE)
    add_arrow(slide, fx+0.85, 4.5, 0.3, gap)
    add_rect(slide, fx-0.3, 4.85, fw+0.6, bh, "1a1a1a", "Bounded Tools", TEAL, 14, TEAL)
    add_arrow(slide, fx+0.85, 5.45, 0.3, gap)
    add_rect(slide, fx, 5.8, fw, bh, "1a1a1a", "Order Draft", AMBER, 14, AMBER)
    add_arrow(slide, fx+0.85, 6.4, 0.3, gap, ORANGE)
    add_rect(slide, fx, 6.75, fw, bh, "1a1a1a", "Human Approval", ORANGE, 14, ORANGE)
    
    # Right: capability panel
    px, pw = 5.0, 7.5
    add_rect(slide, px, 2.0, pw, 2.2, CARD_BG, border=CARD_BORDER)
    add_text(slide, px+0.3, 2.2, pw-0.6, 0.4, "Available to Agent", 16, GREEN, True)
    add_text(slide, px+0.3, 2.7, pw-0.6, 1.4, "✓ search_catalog\n✓ check_inventory\n✓ calculate_order_total\n✓ create_order_draft", 16, GREEN)
    
    add_rect(slide, px, 4.5, pw, 2.2, CARD_BG, border=CARD_BORDER)
    add_text(slide, px+0.3, 4.7, pw-0.6, 0.4, "Not Exposed to Agent", 16, RED, True)
    add_text(slide, px+0.3, 5.2, pw-0.6, 1.4, "✕ approve_order\n✕ reject_order\n✕ deduct_stock\n✕ send_customer_message", 16, RED)
    
    # Bottom
    add_rect(slide, fx, 7.0, fw, bh, "1a1a1a", "Backend", GREEN, 14, GREEN)
    add_rect(slide, fx+2.5, 7.0, fw, bh, "1a1a1a", "Audit Trail", "64748b", 14, "64748b")
    
    if spec.get("notes"):
        add_notes(slide, spec["notes"])
    return slide

def content_slide(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, spec.get("background", BG))
    add_text(slide, 0.8, 0.3, 11.7, 0.8, spec["title"], 40, TITLE, True)
    if spec.get("subtitle"):
        add_text(slide, 0.8, 1.1, 11.7, 0.5, spec["subtitle"], 20, SUBTITLE)
    if spec.get("bullets"):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(spec["bullets"]):
            if isinstance(bullet, str):
                text, bc = bullet, BODY
            else:
                text, bc = bullet.get("text", ""), bullet.get("color", BODY)
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = text
            p.font.size = Pt(22)
            p.font.color.rgb = RGBColor.from_string(bc)
            p.space_after = Pt(16)
    if spec.get("notes"):
        add_notes(slide, spec["notes"])
    return slide

def create_deck(spec_path, output_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for slide_spec in spec["slides"]:
        layout = slide_spec.get("layout", "title_content")
        if layout == "title":
            title_slide(prs, slide_spec)
        elif layout == "architecture":
            arch_slide(prs, slide_spec)
        else:
            content_slide(prs, slide_spec)
    prs.save(output_path)
    print(json.dumps({"ok": True, "output": output_path, "slides": len(prs.slides._sldIdLst)}))

if __name__ == "__main__":
    import sys
    create_deck(sys.argv[1], sys.argv[2])
